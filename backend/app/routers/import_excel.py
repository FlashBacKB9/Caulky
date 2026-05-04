import csv
import io
import re
import uuid
import shutil
from pathlib import Path
from datetime import datetime, date as date_type
from typing import Optional

from fastapi import APIRouter, UploadFile, File, HTTPException, Depends
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl

from app.database import get_db
from app.models.movement import Movement
from app.models.movement_type import MovementType
from app.models.income_expense_group import IncomeExpenseGroup
from app.models.account import Account

router = APIRouter(prefix="/import", tags=["import"])

TEMP_DIR = Path("uploads/import_temp")
TEMP_DIR.mkdir(parents=True, exist_ok=True)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date_type):
        return value.isoformat()
    return str(value)


def _parse_csv(content: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("No se pudo decodificar el archivo CSV")

    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    reader = csv.reader(io.StringIO(text), dialect)
    return [row for row in reader if any(c.strip() for c in row)]


@router.post("/parse")
async def parse_excel(file: UploadFile = File(...)):
    session_id = str(uuid.uuid4())
    filename = (file.filename or "").lower()
    is_csv = filename.endswith(".csv")
    ext = ".csv" if is_csv else ".xlsx"
    temp_path = TEMP_DIR / f"{session_id}{ext}"

    content = await file.read()
    with open(temp_path, "wb") as f:
        f.write(content)

    try:
        if is_csv:
            all_rows = _parse_csv(content)
        else:
            wb = openpyxl.load_workbook(temp_path, data_only=True)
            ws = wb.active
            raw = list(ws.iter_rows(values_only=True))
            wb.close()
            all_rows = [[_cell(c) for c in row] for row in raw]
    except Exception as e:
        temp_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=f"Error al leer el archivo: {e}")

    if not all_rows:
        temp_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="El archivo está vacío")

    columns = [c.strip() or f"Columna {i + 1}" for i, c in enumerate(all_rows[0])]
    rows = [row for row in all_rows[1:] if any(c.strip() for c in row)]

    return {
        "session_id": session_id,
        "columns": columns,
        "rows": rows,
        "total": len(rows),
    }


class TypeMapping(BaseModel):
    action: str  # "auto" | "existing" | "skip" | "create"
    type_id: Optional[int] = None
    group_id: Optional[int] = None


class RunImportRequest(BaseModel):
    session_id: str
    col_date: int
    col_name: int
    col_money: int
    col_bank_date: Optional[int] = None
    col_type: Optional[int] = None
    col_notes: Optional[int] = None
    col_shared: Optional[int] = None
    col_shared_between: Optional[int] = None
    col_my_share: Optional[int] = None
    account_id: Optional[int] = None
    type_map: dict[str, TypeMapping] = {}
    dry_run: bool = False
    skip_duplicates: bool = False


def _parse_money(s: str) -> float:
    s = s.strip().replace(" ", "").replace(" ", "").replace("€", "").replace("$", "").replace("+", "")
    if not s:
        raise ValueError("Importe vacío")
    if "," in s and "." in s:
        if s.rindex(",") > s.rindex("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    return float(s)


_MESES_ES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
    "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
    "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}

def _parse_date(s: str) -> date_type:
    # Normalise: strip, collapse unicode spaces, remove timezone in parens
    s = s.strip()
    s_clean = re.sub(r"\s*\([^)]*\)\s*$", "", s).strip()  # drop "(CET)", "(UTC+1)", …
    s_clean = re.sub(r"\s+", " ", s_clean)                # collapse multiple spaces

    for fmt in [
        "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y",
        "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d",
        "%Y/%m/%d %H:%M", "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
        "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M",
    ]:
        try:
            return datetime.strptime(s_clean, fmt).date()
        except ValueError:
            continue

    # "22 de octubre de 2025"
    match = re.fullmatch(r"(\d{1,2})\s+de\s+(\w+)\s+de\s+(\d{4})", s_clean, re.IGNORECASE)
    if match:
        day = int(match.group(1))
        month_name = match.group(2).lower()
        year = int(match.group(3))
        if month_name in _MESES_ES:
            return date_type(year, _MESES_ES[month_name], day)

    # Last resort: extract first YYYY-MM-DD or YYYY/MM/DD anywhere in the string
    iso = re.search(r"(\d{4})[/\-](\d{1,2})[/\-](\d{1,2})", s_clean)
    if iso:
        return date_type(int(iso.group(1)), int(iso.group(2)), int(iso.group(3)))

    raise ValueError(f"Fecha no reconocida: {s!r}")


@router.post("/run")
async def run_import(req: RunImportRequest, db: AsyncSession = Depends(get_db)):
    temp_path = next(
        (TEMP_DIR / f"{req.session_id}{ext}" for ext in (".xlsx", ".csv")
         if (TEMP_DIR / f"{req.session_id}{ext}").exists()),
        None,
    )
    if temp_path is None:
        raise HTTPException(status_code=404, detail="Sesión expirada, sube el archivo de nuevo")

    # Resolve account: use provided account_id or fall back to main account
    account_id = req.account_id
    if account_id is None:
        res = await db.execute(select(Account).where(Account.is_main == True))
        main_account = res.scalar_one_or_none()
        if main_account:
            account_id = main_account.id

    # Create new movement types (skip on dry_run)
    new_type_ids: dict[str, int] = {}
    if not req.dry_run:
        for type_name, mapping in req.type_map.items():
            if mapping.action == "create" and mapping.group_id:
                grp = await db.get(IncomeExpenseGroup, mapping.group_id)
                mt = MovementType(
                    name=type_name,
                    category=grp.name if grp else "Otros",
                    income_expense_group_id=mapping.group_id,
                )
                db.add(mt)
                await db.flush()
                new_type_ids[type_name] = mt.id

    # Resolve "auto" mappings
    auto_ids: dict[str, int] = {}
    auto_names = [n for n, m in req.type_map.items() if m.action == "auto"]
    if auto_names:
        res = await db.execute(select(MovementType).where(MovementType.name.in_(auto_names)))
        for mt in res.scalars().all():
            auto_ids[mt.name] = mt.id

    # Build existing fingerprints for duplicate detection
    existing: set[tuple] = set()
    if req.skip_duplicates and not req.dry_run:
        res = await db.execute(select(Movement.name, Movement.date, Movement.money))
        existing = {(r[0], str(r[1]), float(r[2])) for r in res.all()}

    try:
        if temp_path.suffix == ".csv":
            data_rows = _parse_csv(temp_path.read_bytes())[1:]
        else:
            wb = openpyxl.load_workbook(temp_path, data_only=True)
            ws = wb.active
            raw = list(ws.iter_rows(values_only=True))
            wb.close()
            data_rows = [[_cell(c) for c in row] for row in raw[1:]]
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error al releer el archivo: {e}")

    imported = 0
    skipped = 0
    errors: list[str] = []
    ok_rows: list[dict] = []

    for i, cells in enumerate(data_rows):

        def get(idx: Optional[int]) -> str:
            if idx is None or idx >= len(cells):
                return ""
            return cells[idx].strip()

        try:
            raw_name = get(req.col_name)
            raw_money = get(req.col_money)
            raw_date = get(req.col_date)
            if not raw_name or not raw_money:
                continue

            money = _parse_money(raw_money)
            parsed_date = _parse_date(raw_date)
            raw_bank_date = get(req.col_bank_date)
            bank_date = _parse_date(raw_bank_date) if raw_bank_date else None
            notes = get(req.col_notes) or None

            raw_shared = get(req.col_shared)
            is_shared = bool(raw_shared) and raw_shared.lower() not in ('0', 'no', 'false')
            shared_between: Optional[int] = None
            my_share: Optional[float] = None
            if is_shared:
                raw_sb = get(req.col_shared_between)
                raw_ms = get(req.col_my_share)
                if raw_ms:
                    try:
                        my_share = _parse_money(raw_ms)
                    except ValueError:
                        pass
                elif raw_sb:
                    try:
                        shared_between = int(float(raw_sb))
                    except ValueError:
                        pass
                if not my_share and not shared_between:
                    shared_between = 2

            type_id: Optional[int] = None
            if req.col_type is not None:
                type_name = get(req.col_type)
                if type_name:
                    mapping = req.type_map.get(type_name)
                    if mapping:
                        if mapping.action == "existing" and mapping.type_id:
                            type_id = mapping.type_id
                        elif mapping.action == "create":
                            type_id = new_type_ids.get(type_name)
                        elif mapping.action == "auto":
                            type_id = auto_ids.get(type_name)

            if req.dry_run:
                ok_rows.append({"name": raw_name, "date": str(parsed_date), "money": money})
            else:
                fingerprint = (raw_name, str(parsed_date), float(money))
                if req.skip_duplicates and fingerprint in existing:
                    skipped += 1
                    continue
                db.add(Movement(
                    name=raw_name,
                    money=money,
                    date=parsed_date,
                    bank_date=bank_date,
                    movement_type_id=type_id,
                    account_id=account_id,
                    paid=True,
                    no_count=False,
                    notes=notes,
                    is_shared=is_shared,
                    shared_between=shared_between,
                    my_share=my_share,
                ))
                imported += 1
        except Exception as e:
            errors.append(f"Fila {i + 2}: {e}")

    if req.dry_run:
        return {"ok": ok_rows, "errors": errors}

    await db.commit()
    temp_path.unlink(missing_ok=True)

    return {"imported": imported, "skipped": skipped, "errors": errors}
